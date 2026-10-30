"""Tests for blank-page detection/removal (PDF and DOCX), the optional
<group> syntax in naming templates, and per-variant settings
(variant_settings in profile.json)."""

import json
import shutil

import pytest
from docx import Document
from pypdf import PdfReader
from reportlab.lib.pagesizes import LETTER
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
from reportlab.lib.styles import getSampleStyleSheet

import generator


def _make_pdf(path, pages):
    """pages: list of strings; "" makes a page with nothing on it."""
    styles = getSampleStyleSheet()
    story = []
    for i, text in enumerate(pages):
        # A page with only a spacer has no text: that is a blank page.
        story.append(Paragraph(text, styles["Normal"]) if text else Spacer(1, 10))
        if i < len(pages) - 1:
            story.append(PageBreak())
    SimpleDocTemplate(str(path), pagesize=LETTER).build(story)
    return path


# --- PDF ---------------------------------------------------------------

class TestPdfBlankPages:
    def test_find_blank_pages(self, tmp_path):
        pdf = _make_pdf(tmp_path / "a.pdf", ["Page one text", "", "Page three text", ""])
        assert generator.find_blank_pdf_pages(pdf) == [1, 3]

    def test_no_blank_pages_returns_empty(self, tmp_path):
        pdf = _make_pdf(tmp_path / "a.pdf", ["One text here", "Two text here"])
        assert generator.find_blank_pdf_pages(pdf) == []
        assert generator.remove_blank_pdf_pages(pdf) == 0

    def test_remove_blank_pages_keeps_content_pages_in_order(self, tmp_path):
        pdf = _make_pdf(tmp_path / "a.pdf", ["First content", "", "Second content", ""])
        assert generator.remove_blank_pdf_pages(pdf) == 2
        reader = PdfReader(str(pdf))
        assert len(reader.pages) == 2
        assert "First content" in reader.pages[0].extract_text()
        assert "Second content" in reader.pages[1].extract_text()
        assert not Path_exists_tmp(pdf)

    def test_never_empties_an_all_blank_pdf(self, tmp_path, capsys):
        pdf = _make_pdf(tmp_path / "a.pdf", ["", ""])
        assert generator.remove_blank_pdf_pages(pdf) == 0
        assert len(PdfReader(str(pdf)).pages) == 2
        assert "looks blank" in capsys.readouterr().err

    def test_render_resume_pdf_removes_trailing_blank_pages(self, tmp_path, monkeypatch, sample_resume_dict):
        real_story = generator._build_resume_story
        monkeypatch.setattr(
            generator, "_build_resume_story",
            lambda r, tier: real_story(r, tier) + [PageBreak(), PageBreak()],
        )
        pdf = tmp_path / "resume.pdf"
        pages, body_pt = generator.render_resume_pdf(sample_resume_dict, pdf)
        assert pages == 1
        assert len(PdfReader(str(pdf)).pages) == 1
        assert generator.find_blank_pdf_pages(pdf) == []

    def test_blank_page_never_shrinks_the_font(self, tmp_path, monkeypatch, sample_resume_dict):
        """Without blank-page removal, 2 extra blank pages would push the
        page count past max_pages and force the smallest font tier."""
        real_story = generator._build_resume_story
        monkeypatch.setattr(
            generator, "_build_resume_story",
            lambda r, tier: real_story(r, tier) + [PageBreak(), PageBreak(), PageBreak()],
        )
        _, body_pt = generator.render_resume_pdf(sample_resume_dict, tmp_path / "resume.pdf")
        assert body_pt == generator.PDF_TIERS[0][0]

    def test_cover_letter_pdf_has_no_blank_pages(self, tmp_path, sample_cover_letter_dict):
        cl = dict(sample_cover_letter_dict, body=sample_cover_letter_dict["body"] + "\n\n\n\n\n\n")
        pdf = tmp_path / "cl.pdf"
        generator.render_cover_letter_pdf(cl, pdf)
        assert generator.find_blank_pdf_pages(pdf) == []


def Path_exists_tmp(pdf):
    from pathlib import Path
    return Path(str(pdf) + ".tmp").exists()


# --- DOCX --------------------------------------------------------------

class TestDocxBlankPageSources:
    def test_removes_manual_page_breaks_and_empty_paragraphs(self):
        from docx.enum.text import WD_BREAK
        doc = Document()
        doc.add_paragraph("Real content")
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_paragraph("")
        doc.add_paragraph("More content")
        doc.add_paragraph("   ")
        removed = generator.strip_docx_blank_page_sources(doc)
        assert removed >= 3
        assert [p.text for p in doc.paragraphs] == ["Real content", "More content"]
        assert 'w:type="page"' not in doc.element.xml

    def test_keeps_empty_paragraph_after_a_final_table(self):
        doc = Document()
        doc.add_paragraph("Intro")
        doc.add_table(rows=1, cols=1)
        doc.add_paragraph("")
        generator.strip_docx_blank_page_sources(doc)
        children = [c.tag.split("}")[1] for c in doc.element.body]
        assert children[-2:] == ["tbl", "p"] or children[-3:-1] == ["tbl", "p"]

    def test_keeps_paragraph_between_two_tables(self):
        doc = Document()
        doc.add_table(rows=1, cols=1)
        doc.add_paragraph("")
        doc.add_table(rows=1, cols=1)
        generator.strip_docx_blank_page_sources(doc)
        assert len(doc.paragraphs) == 1

    def test_keeps_paragraph_containing_a_drawing(self):
        doc = Document()
        p = doc.add_paragraph("")
        from docx.oxml import parse_xml
        p._p.append(parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:drawing/></w:r>'))
        generator.strip_docx_blank_page_sources(doc)
        assert len(doc.paragraphs) == 1

    def test_rendered_resume_docx_has_no_empty_body_paragraphs(self, tmp_path, sample_resume_dict):
        path = tmp_path / "resume.docx"
        generator.render_resume_docx(sample_resume_dict, path, body_pt=10.5)
        doc = Document(str(path))
        assert all(p.text.strip() for p in doc.paragraphs)

    def test_cover_letter_docx_skips_empty_paragraphs(self, tmp_path, sample_cover_letter_dict):
        cl = dict(sample_cover_letter_dict, body="One.\n\n\n\n\n\nTwo.")
        path = tmp_path / "cl.docx"
        generator.render_cover_letter_docx(cl, path)
        assert [p.text for p in Document(str(path)).paragraphs] == ["One.", "Two."]


needs_soffice = pytest.mark.skipif(
    shutil.which("soffice") is None and shutil.which("libreoffice") is None,
    reason="LibreOffice not installed",
)


class TestDocxRenderedCheck:
    def test_find_blank_docx_pages_returns_none_without_libreoffice(self, tmp_path, monkeypatch):
        monkeypatch.setattr(generator.shutil, "which", lambda name: None)
        assert generator.find_blank_docx_pages(tmp_path / "x.docx") is None

    def test_report_says_skipped_without_libreoffice(self, tmp_path, monkeypatch, capsys):
        monkeypatch.setattr(generator.shutil, "which", lambda name: None)
        generator.report_docx_blank_pages(tmp_path / "x.docx")
        assert "Skipped" in capsys.readouterr().out

    @needs_soffice
    def test_rendered_check_finds_then_stripping_removes_blank_pages(self, tmp_path):
        from docx.enum.text import WD_BREAK
        doc = Document()
        doc.add_paragraph("Page one")
        for _ in range(2):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_paragraph("Page four")
        before = tmp_path / "before.docx"
        doc.save(before)
        assert generator.find_blank_docx_pages(before)

        generator.strip_docx_blank_page_sources(doc)
        after = tmp_path / "after.docx"
        doc.save(after)
        assert generator.find_blank_docx_pages(after) == []


# --- <optional groups> in naming templates -----------------------------

TEMPLATE = "<{JobAcronym}/>{FirstName} {LastName} Resume< ({JobAcronym})>.{Extension}"


class TestOptionalGroups:
    def test_groups_kept_when_acronym_present(self):
        assert generator.render_filename(TEMPLATE, "Jane Doe", "SDE", "pdf") == "SDE/Jane Doe Resume (SDE).pdf"

    @pytest.mark.parametrize("blank", ["", "   ", None])
    def test_groups_dropped_when_acronym_blank(self, blank):
        assert generator.render_filename(TEMPLATE, "Jane Doe", blank, "pdf") == "Jane Doe Resume.pdf"

    def test_template_without_groups_is_unchanged(self):
        assert generator.render_filename(
            "{FirstName} {LastName} Resume ({JobAcronym}).{Extension}", "Jane Doe", "SDE", "pdf"
        ) == "Jane Doe Resume (SDE).pdf"

    def test_group_with_only_literal_text_is_kept(self):
        assert generator.render_filename("A<-literal>.{Extension}", "", "", "txt") == "A-literal.txt"

    def test_group_with_format_spec_placeholder(self):
        out = generator.render_filename("<{datetime.now.year:04d}->x", "", "", "")
        assert out.endswith("-x") and out[:4].isdigit()

    def test_unmatched_angle_bracket_is_left_alone(self):
        assert generator.render_filename("a<b", "", "", "") == "a<b"

    def test_blank_acronym_does_not_leak_other_groups(self):
        out = generator.render_filename("<{Email}_>{FirstName}.{Extension}", "Jane Doe", "", "pdf", email="j@x.com")
        assert out == "j@x.com_Jane.pdf"


# --- variant_settings --------------------------------------------------

class TestVariantSettings:
    def test_acronym_defaults_to_variant_name(self, sample_kb):
        assert generator.resolve_job_acronym(sample_kb, "SDE") == "SDE"

    @pytest.mark.parametrize("override,expected", [("", ""), (None, ""), ("   ", ""), (" GEN ", "GEN")])
    def test_acronym_override(self, sample_kb, override, expected):
        sample_kb["variant_settings"] = {"GENERIC": {"job_acronym": override}}
        assert generator.resolve_job_acronym(sample_kb, "GENERIC") == expected

    def test_settings_without_job_acronym_key_keep_variant_name(self, sample_kb):
        sample_kb["variant_settings"] = {"GENERIC": {"skills_heading": "X"}}
        assert generator.resolve_job_acronym(sample_kb, "GENERIC") == "GENERIC"

    def test_non_dict_settings_entry_is_ignored(self, sample_kb):
        sample_kb["variant_settings"] = {"_note": "just a string", "GENERIC": "oops"}
        assert generator.get_variant_settings(sample_kb, "GENERIC") == {}
        assert generator.get_variant_settings(sample_kb, "_note") == {}

    def test_skills_heading_override(self, sample_kb):
        sample_kb["variant_settings"] = {"GENERIC": {"skills_heading": "CORE TECHNICAL SKILLS"}}
        ctx = generator.build_baseline_context(sample_kb, "GENERIC")
        assert ctx["skills_heading"] == "CORE TECHNICAL SKILLS"

    def test_skills_heading_without_override_unchanged(self, sample_kb):
        assert generator.build_baseline_context(sample_kb, "SRE")["skills_heading"] == "CORE SRE SKILLS"


class TestMainGenericVariant:
    @pytest.fixture
    def generic_env(self, tmp_path, monkeypatch, sample_kb, sample_resume_dict, sample_cover_letter_dict):
        sample_kb["summary_variants"]["GENERIC"] = "Generic summary."
        sample_kb["variant_settings"] = {"GENERIC": {"job_acronym": "", "skills_heading": "CORE TECHNICAL SKILLS"}}
        monkeypatch.chdir(tmp_path)
        (tmp_path / "data").mkdir()
        (tmp_path / "data" / "profile.json").write_text(json.dumps(sample_kb), encoding="utf-8")
        (tmp_path / "RESUME.template.md").write_text("t", encoding="utf-8")
        (tmp_path / "README.template.md").write_text("t", encoding="utf-8")
        monkeypatch.setenv("LITELLM_BASE_URL", "http://example.com")
        monkeypatch.setenv("LITELLM_API_KEY", "secret")
        monkeypatch.setenv("VARIANTS", '["GENERIC", "SDE"]')
        monkeypatch.setenv("RESUME_NAMING_TEMPLATE", TEMPLATE)
        monkeypatch.setenv("COVERLETTER_NAMING_TEMPLATE",
                           "<{JobAcronym}/>{FirstName} {LastName} Cover Letter< ({JobAcronym})>.{Extension}")
        monkeypatch.delenv("OUTPUT_FOLDER", raising=False)
        monkeypatch.delenv("DATA", raising=False)
        monkeypatch.setattr(generator, "call_llm_fill_resume", lambda c, kb, v, t: sample_resume_dict)
        monkeypatch.setattr(generator, "call_llm_cover_letter", lambda c, kb, v: sample_cover_letter_dict)
        monkeypatch.setattr(generator, "call_llm_readme", lambda c, kb, t: "# Jane\n\nbody")
        return tmp_path

    def test_generic_files_have_no_acronym_or_subfolder(self, generic_env):
        generator.main()
        out = generic_env / "generated"
        for ext in ("json", "txt", "md", "pdf", "docx"):
            assert (out / f"Jane Doe Resume.{ext}").exists()
        for ext in ("txt", "docx", "pdf"):
            assert (out / f"Jane Doe Cover Letter.{ext}").exists()

    def test_other_variants_keep_subfolder_and_acronym(self, generic_env):
        generator.main()
        out = generic_env / "generated"
        assert (out / "SDE" / "Jane Doe Resume (SDE).pdf").exists()
        assert (out / "SDE" / "Jane Doe Cover Letter (SDE).docx").exists()

    def test_no_pdf_in_output_has_blank_pages(self, generic_env):
        generator.main()
        for pdf in (generic_env / "generated").rglob("*.pdf"):
            assert generator.find_blank_pdf_pages(pdf) == [], pdf
